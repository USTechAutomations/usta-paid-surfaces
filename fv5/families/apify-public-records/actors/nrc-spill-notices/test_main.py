import asyncio
import io
import math
import time
import unittest
from datetime import date
from unittest.mock import patch
from openpyxl import Workbook
import main

COMMON = ['SEQNOS','INCIDENT_DATE_TIME','INCIDENT_DTG','LOCATION_STATE','LOCATION_NEAREST_CITY','TYPE_OF_INCIDENT']
MATERIAL = ['SEQNOS','NAME_OF_MATERIAL','AMOUNT_OF_MATERIAL','UNIT_OF_MEASURE']
DETAIL = ['SEQNOS','MEDIUM_DESC']


def fixture(*, common=None, materials=None, details=None, missing=None):
    w=Workbook();w.remove(w.active)
    for name,head,rows in [('INCIDENT_COMMONS',COMMON,common if common is not None else [[10,'1/1/2026 01:00','OCCURRED','LA','TEST CITY','VESSEL'],[11,'1/2/2026 01:00','DISCOVERED','TX','TEST CITY','FIXED']]),('MATERIAL_INVOLVED',MATERIAL,materials if materials is not None else [[10,'OIL',0,'UNKNOWN AMOUNT'],[10,'DIESEL',5,'GALLON(S)'],[11,'AMMONIA',10,'POUND(S)']]),('INCIDENT_DETAILS',DETAIL,details if details is not None else [[10,'WATER'],[10,'LAND'],[11,'AIR']])]:
        if name==missing:continue
        s=w.create_sheet(name);s.append(head)
        for row in rows:s.append(row)
    s=w.create_sheet('CALLS');s.append(['SEQNOS','CALLER_NAME','PHONE','ADDRESS']);s.append([10,'SYNTHETIC PRIVATE PERSON','555-0100','123 TEST ST'])
    out=io.BytesIO();w.save(out);return out.getvalue()


class InputTests(unittest.TestCase):
    def test_small_default_and_receipt_year(self):
        x=main.validate_input({},today=date(2026,9,12));self.assertEqual((x['maxItems'],x['year']),(10,2026))
    def test_boundaries_and_types(self):
        for inp in [{'maxItems':True},{'maxItems':0},{'maxItems':1001},{'maxItems':'10'},{'year':2030},{'year':1989},{'timeoutSeconds':1},{'state':[]},{'state':'ZZ'},{'dateFrom':'2026-02-30'},{'dateFrom':'2026-1-1'},{'other':1},[],{'materialKeyword':'a'*101}]:
            with self.subTest(inp=inp),self.assertRaises(ValueError):main.validate_input(inp,today=date(2026,9,12))
    def test_cross_year_requires_receipt_year(self):
        with self.assertRaises(ValueError):main.validate_input({'dateFrom':'2025-12-31','dateTo':'2026-01-02'})
        x=main.validate_input({'year':2026,'dateFrom':'2025-12-31','dateTo':'2026-01-02'});self.assertEqual(x['year'],2026)
    def test_reversed_dates(self):
        with self.assertRaises(ValueError):main.validate_input({'dateFrom':'2026-02-01','dateTo':'2026-01-01'})
    def test_invalid_input_never_fetches(self):
        with patch.object(main,'fetch_workbook') as fetch,self.assertRaises(ValueError):main.collect({'year':True})
        fetch.assert_not_called()


class WorkbookTests(unittest.TestCase):
    def test_join_preserves_one_report_multiple_materials_and_units(self):
        items,summary=main.collect({'year':2026},body=fixture());self.assertEqual(len(items),2)
        x=next(i for i in items if i['report_id']=='10');self.assertEqual(len(x['materials']),2);self.assertIsNone(x['quantity']);self.assertEqual(x['medium'],'LAND; WATER');self.assertEqual(summary['matched_reports'],2)
    def test_unknown_amount_is_not_zero(self):
        items,_=main.collect({'year':2026,'state':'LA'},body=fixture());m=items[0]['materials'][0]
        self.assertIsNone(m['quantity']);self.assertEqual((m['source_amount'],m['quantity_status']),(0,'UNKNOWN'))
    def test_zero_with_known_unit_is_preserved(self):
        items,_=main.collect({'year':2026,'state':'LA'},body=fixture(materials=[[10,'OIL',0,'GALLON(S)']]))
        self.assertEqual(items[0]['quantity'],0);self.assertEqual(items[0]['materials'][0]['quantity_status'],'REPORTED')
    def test_dates_numeric_sort_and_filter(self):
        b=fixture(common=[[10,'12/31/2025 01:00','DISCOVERED','LA','TEST','VESSEL'],[11,'1/1/2026 01:00','OCCURRED','TX','TEST','FIXED']])
        items,_=main.collect({'year':2026,'dateFrom':'2026-01-01','dateTo':'2026-01-01'},body=b);self.assertEqual([x['report_id'] for x in items],['11'])
    def test_malformed_date_unknown_preserved_excluded_only_by_date_filter(self):
        b=fixture(common=[[10,'6/11/205 20:25','OCCURRED','LA','TEST','VESSEL'] ],materials=[[10,'OIL',5,'GALLON(S)']],details=[[10,'WATER']])
        items,_=main.collect({'year':2026},body=b);self.assertIsNone(items[0]['incident_date']);self.assertEqual(items[0]['date_status'],'UNKNOWN');self.assertEqual(items[0]['incident_date_source'],'6/11/205 20:25')
        items,s=main.collect({'year':2026,'dateFrom':'2026-01-01'},body=b);self.assertEqual(items,[]);self.assertEqual(s['missing_dates_excluded_by_filter'],1)
    def test_non_material_reports_and_private_columns_not_emitted(self):
        items,s=main.collect({'year':2026},body=fixture(materials=[[10,'OIL',1,'GALLON(S)']]))
        self.assertEqual(len(items),1);self.assertNotIn('SYNTHETIC PRIVATE PERSON',str(items));self.assertNotIn('123 TEST ST',str(items));self.assertEqual(s['source_incident_rows'],2)
    def test_keyword_only_matches_material(self):
        items,s=main.collect({'year':2026,'materialKeyword':'diesel'},body=fixture());self.assertEqual(items[0]['report_id'],'10')
        items,s=main.collect({'year':2026,'materialKeyword':'TEST CITY'},body=fixture());self.assertEqual(items,[]);self.assertEqual(s['status'],'NO_MATCHES')
    def test_cap_reports_omitted_count(self):
        items,s=main.collect({'year':2026,'maxItems':1},body=fixture());self.assertEqual(items[0]['report_id'],'11');self.assertEqual(s['omitted_due_to_max_items'],1)
    def test_html_corrupt_and_missing_sheet_fail(self):
        for b in [b'<html>Application error</html>',b'PK\x03\x04oops',fixture(missing='MATERIAL_INVOLVED')]:
            with self.subTest(bodylen=len(b)),self.assertRaises(main.SourceUnavailable):main.collect({'year':2026},body=b)
    def test_empty_incident_source_unknown(self):
        with self.assertRaises(main.SourceUnavailable):main.collect({'year':2026},body=fixture(common=[],materials=[],details=[]))
    def test_duplicate_and_orphan_ids_fail(self):
        for b in [fixture(common=[[10,'1/1/2026','OCCURRED','LA','TEST','VESSEL']]*2),fixture(materials=[[999,'OIL',5,'GALLON(S)']])]:
            with self.assertRaises(main.SourceUnavailable):main.collect({'year':2026},body=b)
    def test_unjoined_narrative_rows_counted_without_exposing_text(self):
        items,s=main.collect({'year':2026},body=fixture(details=[[10,'WATER'],['SYNTHETIC CONTINUATION WITH PRIVATE TEXT',None]]))
        self.assertEqual(s['unjoined_details_rows'],1);self.assertNotIn('PRIVATE TEXT',str((items,s)))
    def test_unjoined_medium_fails_instead_of_guessing_report(self):
        with self.assertRaises(main.SourceUnavailable):main.collect({'year':2026},body=fixture(details=[['BAD ID','WATER']]))
    def test_invalid_quantity_fails(self):
        with self.assertRaises(main.SourceUnavailable):main.collect({'year':2026},body=fixture(materials=[[10,'OIL',-1,'GALLON(S)']]))
    def test_timeout_fails_instead_of_partial_success(self):
        with self.assertRaises(main.SourceUnavailable):main.parse_workbook(fixture(),main.validate_input({'year':2026}),time.monotonic()-1)
    def test_size_limit_fails(self):
        with patch.object(main,'MAX_BYTES',1),self.assertRaises(main.SourceUnavailable):main.collect({'year':2026},body=fixture())
    def test_provenance_and_report_status(self):
        items,s=main.collect({'year':2026},body=fixture());self.assertEqual(items[0]['source_url'],'https://nrc.uscg.mil/FOIAFiles/CY26.xlsx');self.assertEqual(items[0]['report_status'],'INITIAL_UNVALIDATED');self.assertEqual(len(s['source_sha256']),64)


class FetchTests(unittest.TestCase):
    def test_http_error_unknown(self):
        with patch.object(main.urllib.request,'urlopen',side_effect=OSError('offline')),self.assertRaises(main.SourceUnavailable):main.fetch_workbook(main.source_url(2026),time.monotonic()+40)
    def test_redirect_and_html_unknown(self):
        class Reply(io.BytesIO):
            status=200
        for url,body in [('https://nrc.uscg.mil/ApplicationError.aspx',b'PK\x03\x04'),(main.source_url(2026),b'<html>error</html>')]:
            reply=Reply(body);reply.url=url
            with patch.object(main.urllib.request,'urlopen',return_value=reply),self.assertRaises(main.SourceUnavailable):main.fetch_workbook(main.source_url(2026),time.monotonic()+40)


if __name__=='__main__':unittest.main()
