#!/usr/bin/env python3
"""NRC annual initial-report workbook, one item per report with listed material.

Only official USCG workbook URLs are used. Reports are unvalidated; no caller,
address, narrative, vehicle, or responsible-party fields are emitted. Source
errors fail the run before any dataset output. Apify's existing automatic
per-dataset-item event owns billing; there are no manual charge calls.
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import math
import re
import time
import urllib.request
import zipfile
from datetime import date, datetime, timezone
from openpyxl import load_workbook

BASE = 'https://nrc.uscg.mil'
MAX_BYTES = 64 * 1024 * 1024
MAX_EXPANDED_BYTES = 512 * 1024 * 1024
MAX_ROWS = 300_000
STATES = set('AL AK AZ AR CA CO CT DE DC FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY AS GU MP PR VI'.split())
UA = 'USTechAutomations-NRC/0.3 (+https://ustechautomations.com/feeds)'


class SourceUnavailable(RuntimeError):
    """The source cannot support a result; absence of reports is UNKNOWN."""


def _remaining(deadline):
    value = deadline - time.monotonic()
    if value <= 0:
        raise SourceUnavailable('UNKNOWN: source processing exceeded timeout; no results published')
    return value


def _text(value):
    return '' if value is None else str(value).strip()


def _date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    if not text or text.upper() in {'UNKNOWN', 'NOT PROVIDED', 'N/A'}:
        return None
    for fmt in ('%m/%d/%Y %H:%M', '%m/%d/%Y %H:%M:%S', '%m/%d/%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None  # Preserve malformed source dates as UNKNOWN; never invent a correction.


def validate_input(inp, today=None):
    if not isinstance(inp, dict):
        raise ValueError('Input must be an object')
    allowed = {'year', 'state', 'dateFrom', 'dateTo', 'materialKeyword', 'keyword', 'maxItems', 'timeoutSeconds'}
    if set(inp) - allowed:
        raise ValueError('Unknown input fields: ' + ', '.join(sorted(set(inp) - allowed)))
    today = today or datetime.now(timezone.utc).date()
    out = {}
    for key in ('state', 'dateFrom', 'dateTo', 'materialKeyword', 'keyword'):
        val = inp.get(key, '')
        if not isinstance(val, str):
            raise ValueError(key + ' must be text')
        out[key] = val.strip()
    out['state'] = out['state'].upper()
    if out['state'] and out['state'] not in STATES:
        raise ValueError('state must be a US state or territory postal abbreviation')
    for key in ('dateFrom', 'dateTo'):
        if out[key]:
            if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', out[key]):
                raise ValueError(key + ' must use YYYY-MM-DD')
            date.fromisoformat(out[key])
    if out['dateFrom'] and out['dateTo'] and out['dateFrom'] > out['dateTo']:
        raise ValueError('dateFrom must not be after dateTo')
    # Annual files use report receipt year, not necessarily incident year.
    if 'year' not in inp and out['dateFrom'] and out['dateTo'] and out['dateFrom'][:4] != out['dateTo'][:4]:
        raise ValueError('Select an explicit report receipt year for a cross-year incident-date range')
    default_year = 2026  # Only this annual workbook has current source acceptance.
    for key, default, low, high in [('year', default_year, 2026, 2026), ('maxItems', 10, 1, 1000), ('timeoutSeconds', 180, 30, 600)]:
        val = inp.get(key, default)
        if type(val) is not int or not low <= val <= high:
            raise ValueError(f'{key} must be an integer from {low} to {high}')
        out[key] = val
    if len(out['materialKeyword']) > 100 or len(out['keyword']) > 100:
        raise ValueError('Material keyword must be at most 100 characters')
    if out['keyword'] and out['materialKeyword'] and out['keyword'] != out['materialKeyword']:
        raise ValueError('Provide only one material keyword')
    out['materialKeyword'] = out['materialKeyword'] or out['keyword']
    return out


def source_url(year):
    return f'{BASE}/FOIAFiles/CY{year % 100:02d}.xlsx'


def fetch_workbook(url, deadline):
    try:
        request = urllib.request.Request(url, headers={'User-Agent': UA})
        with urllib.request.urlopen(request, timeout=min(45, _remaining(deadline))) as response:
            if response.status != 200 or response.url != url:
                raise SourceUnavailable('UNKNOWN: workbook request did not return the requested official file')
            chunks, count = [], 0
            while chunk := response.read(1024 * 1024):
                _remaining(deadline)
                count += len(chunk)
                if count > MAX_BYTES:
                    raise SourceUnavailable('UNKNOWN: source file exceeds supported size')
                chunks.append(chunk)
            body = b''.join(chunks)
    except SourceUnavailable:
        raise
    except Exception as exc:
        raise SourceUnavailable('UNKNOWN: official workbook could not be downloaded') from exc
    if not body.startswith(b'PK\x03\x04'):
        raise SourceUnavailable('UNKNOWN: source returned an error page or non-workbook content')
    return body


def _table(workbook, sheet, required, deadline, diagnostics=None):
    if sheet not in workbook.sheetnames:
        raise SourceUnavailable('UNKNOWN: missing source worksheet ' + sheet)
    rows = workbook[sheet].iter_rows(values_only=True)
    header = next(rows, None)
    if not header or len(set(header)) != len(header) or not set(required).issubset(header):
        raise SourceUnavailable('UNKNOWN: source columns changed in ' + sheet)
    indexes = {name: header.index(name) for name in required}
    for n, row in enumerate(rows, 1):
        _remaining(deadline)
        if n > MAX_ROWS:
            raise SourceUnavailable('UNKNOWN: source sheet exceeds supported row count')
        if not any(v is not None for v in row):
            continue
        item = {name: row[index] if index < len(row) else None for name, index in indexes.items()}
        raw_id = item['SEQNOS']
        if isinstance(raw_id, bool) or not re.fullmatch(r'[1-9]\d*', _text(raw_id)):
            if sheet == 'INCIDENT_DETAILS' and not _text(item.get('MEDIUM_DESC')) and diagnostics is not None:
                diagnostics['unjoined_details_rows'] += 1
                continue  # Official sheet contains narrative continuation rows with no medium.
            raise SourceUnavailable('UNKNOWN: source report identifier is missing or malformed')
        item['SEQNOS'] = _text(raw_id)
        yield item


def parse_workbook(body, inp, deadline):
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            if sum(x.file_size for x in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise SourceUnavailable('UNKNOWN: expanded workbook exceeds supported size')
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except SourceUnavailable:
        raise
    except Exception as exc:
        raise SourceUnavailable('UNKNOWN: source workbook cannot be read') from exc
    try:
        reports = {}
        for row in _table(workbook, 'INCIDENT_COMMONS', ['SEQNOS', 'INCIDENT_DATE_TIME', 'INCIDENT_DTG', 'LOCATION_STATE', 'LOCATION_NEAREST_CITY', 'TYPE_OF_INCIDENT'], deadline):
            rid = row['SEQNOS']
            if rid in reports:
                raise SourceUnavailable('UNKNOWN: duplicate incident identifier')
            reports[rid] = {'report_id': rid, 'incident_date': _date(row['INCIDENT_DATE_TIME']),
                'incident_date_source': _text(row['INCIDENT_DATE_TIME']) or None,
                'date_status': 'REPORTED' if _date(row['INCIDENT_DATE_TIME']) else 'UNKNOWN',
                'date_basis': _text(row['INCIDENT_DTG']) or None, 'state': _text(row['LOCATION_STATE']),
                'city': _text(row['LOCATION_NEAREST_CITY']), 'incident_type': _text(row['TYPE_OF_INCIDENT']),
                'materials': [], 'medium': []}
        if not reports:
            raise SourceUnavailable('UNKNOWN: source contains no incident records')
        material_rows = 0
        for row in _table(workbook, 'MATERIAL_INVOLVED', ['SEQNOS', 'NAME_OF_MATERIAL', 'AMOUNT_OF_MATERIAL', 'UNIT_OF_MEASURE'], deadline):
            material_rows += 1
            if row['SEQNOS'] not in reports:
                raise SourceUnavailable('UNKNOWN: material refers to a missing incident')
            amount, unit = row['AMOUNT_OF_MATERIAL'], _text(row['UNIT_OF_MEASURE'])
            if amount is not None and (type(amount) not in (int, float) or not math.isfinite(amount) or amount < 0):
                raise SourceUnavailable('UNKNOWN: invalid material quantity')
            known = amount is not None and bool(unit) and 'UNKNOWN' not in unit.upper() and unit.upper() not in {'N/A', 'NOT PROVIDED'}
            material = {'name': _text(row['NAME_OF_MATERIAL']) or None,
                        'quantity': amount if known else None, 'unit': unit or None,
                        'quantity_status': 'REPORTED' if known else 'UNKNOWN', 'source_amount': amount}
            materials = reports[row['SEQNOS']]['materials']
            if material not in materials:
                materials.append(material)
        diagnostics = {'unjoined_details_rows': 0}
        for row in _table(workbook, 'INCIDENT_DETAILS', ['SEQNOS', 'MEDIUM_DESC'], deadline, diagnostics):
            if row['SEQNOS'] not in reports:
                raise SourceUnavailable('UNKNOWN: medium refers to a missing incident')
            medium = _text(row['MEDIUM_DESC'])
            values = reports[row['SEQNOS']]['medium']
            if medium and medium not in values:
                values.append(medium)
        matches = []
        missing_date_excluded = 0
        for item in reports.values():
            _remaining(deadline)
            if not item['materials']:
                continue
            if inp['state'] and item['state'].upper() != inp['state']:
                continue
            if inp['materialKeyword'] and not any(inp['materialKeyword'].casefold() in (m['name'] or '').casefold() for m in item['materials']):
                continue
            d = item['incident_date']
            if (inp['dateFrom'] or inp['dateTo']) and d is None:
                missing_date_excluded += 1
                continue
            if d and ((inp['dateFrom'] and d < inp['dateFrom']) or (inp['dateTo'] and d > inp['dateTo'])):
                continue
            item['material'] = '; '.join(m['name'] or 'UNKNOWN' for m in item['materials'])
            item['quantity'] = item['materials'][0]['quantity'] if len(item['materials']) == 1 else None
            item['quantity_unit'] = item['materials'][0]['unit'] if len(item['materials']) == 1 else None
            item['medium'] = '; '.join(sorted(item['medium'])) or None
            item['medium_status'] = 'REPORTED' if item['medium'] and 'UNKNOWN' not in item['medium'] else 'UNKNOWN'
            matches.append(item)
        matches.sort(key=lambda item: (item['incident_date'] or '', int(item['report_id'])), reverse=True)
        return matches[:inp['maxItems']], {**diagnostics, 'source_incident_rows': len(reports), 'source_material_rows': material_rows,
            'matched_reports': len(matches), 'omitted_due_to_max_items': max(0, len(matches) - inp['maxItems']),
            'missing_dates_excluded_by_filter': missing_date_excluded}
    except SourceUnavailable:
        raise
    except Exception as exc:
        raise SourceUnavailable('UNKNOWN: workbook content could not be processed') from exc
    finally:
        workbook.close()


def collect(inp, *, body=None, fetch=None):
    values = validate_input(inp)
    deadline = time.monotonic() + values['timeoutSeconds']
    url = source_url(values['year'])
    body = body if body is not None else (fetch or fetch_workbook)(url, deadline)
    if len(body) > MAX_BYTES:
        raise SourceUnavailable('UNKNOWN: source file exceeds supported size')
    items, counts = parse_workbook(body, values, deadline)
    provenance = {'source_url': url, 'source_sha256': hashlib.sha256(body).hexdigest(),
        'source_retrieved_at': datetime.now(timezone.utc).isoformat(), 'report_receipt_year': values['year'],
        'report_status': 'INITIAL_UNVALIDATED', 'source_scope': 'MATERIAL_INVOLVED worksheet; continuous-release-only and reports without a listed material excluded'}
    for item in items:
        item.update(provenance)
    return items, {'status': 'SUCCESS' if items else 'NO_MATCHES', **counts, **provenance,
        'returned_reports': len(items), 'filters': values}


async def main():
    from apify import Actor
    async with Actor:
        try:
            raw_input = await Actor.get_input()
            items, summary = await asyncio.to_thread(collect, {} if raw_input is None else raw_input)
        except Exception as exc:
            await Actor.set_value('SUMMARY', {'status': 'UNKNOWN' if isinstance(exc, SourceUnavailable) else 'INVALID_INPUT', 'error': str(exc), 'returned_reports': None})
            raise
        final_status = summary['status']
        summary.update({'status': 'OUTPUT_IN_PROGRESS', 'prepared_reports': len(items), 'returned_reports': None})
        await Actor.set_value('SUMMARY', summary)
        # Dataset-item pricing is applied automatically by Apify. Never double-charge.
        for item in items:
            await Actor.push_data(item)
        summary.update({'status': final_status, 'returned_reports': len(items)})
        await Actor.set_value('SUMMARY', summary)
        Actor.log.info(f"Source {summary['status']}: {len(items)} reports returned; {summary['omitted_due_to_max_items']} matches omitted by maxItems")


if __name__ == '__main__':
    asyncio.run(main())
